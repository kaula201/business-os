"""GL enhanced: period closing, analytics, export, require_module migration."""
from datetime import date, datetime
from decimal import Decimal
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.core.dependencies import get_current_user, require_module
from app.models.user import User
from app.models.gl import GLAccount, JournalEntry, JournalEntryLine
from app.models.accounting_period import AccountingPeriod
from app.schemas.common import ResponseBase
from pydantic import BaseModel, Field
from io import BytesIO
import openpyxl
from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/gl", tags=["მთავარი წიგნი — გაძლიერებული"])


# ── Period Closing ────────────────────────────────────────────────────────────

class PeriodCloseRequest(BaseModel):
    period_id: UUID


class PeriodCloseResponse(BaseModel):
    period_id: UUID
    period_label: str
    status: str
    entries_count: int
    message: str


@router.post("/period-close", response_model=ResponseBase[PeriodCloseResponse])
async def close_accounting_period(
    data: PeriodCloseRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_module("gl", "can_approve")),
):
    period = (await db.execute(
        select(AccountingPeriod).where(
            AccountingPeriod.id == data.period_id,
            AccountingPeriod.company_id == current_user.company_id,
        ).with_for_update()
    )).scalar_one_or_none()
    if not period:
        raise HTTPException(status_code=404, detail="სააღრიცხვო პერიოდი არ მოიძებნა")
    if period.status == "closed":
        raise HTTPException(status_code=409, detail="პერიოდი უკვე დახურულია")

    # Count entries in this period
    entries_count = (await db.execute(
        select(func.count(JournalEntry.id)).where(
            JournalEntry.company_id == current_user.company_id,
            JournalEntry.entry_date >= period.start_date,
            JournalEntry.entry_date <= period.end_date,
        )
    )).scalar()

    period.status = "closed"
    await db.flush()

    return ResponseBase(data=PeriodCloseResponse(
        period_id=period.id,
        period_label=f"{period.start_date} — {period.end_date}",
        status="closed",
        entries_count=entries_count or 0,
        message=f"სააღრიცხვო პერიოდი დახურულია. {entries_count} ჩანაწერი დაფიქსირდა.",
    ))


# ── GL Analytics ──────────────────────────────────────────────────────────────

class GLAnalytics(BaseModel):
    total_accounts: int
    active_accounts: int
    total_journal_entries: int
    entries_this_month: int
    total_debit: Decimal
    total_credit: Decimal
    period_status: str | None


@router.get("/analytics", response_model=ResponseBase[GLAnalytics])
async def gl_analytics(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_module("gl", "can_access")),
):
    company_id = current_user.company_id
    today = date.today()
    month_start = today.replace(day=1)

    total_accounts = (await db.execute(
        select(func.count(GLAccount.id)).where(GLAccount.company_id == company_id)
    )).scalar()
    active_accounts = (await db.execute(
        select(func.count(GLAccount.id)).where(GLAccount.company_id == company_id, GLAccount.is_active == True)
    )).scalar()

    total_entries = (await db.execute(
        select(func.count(JournalEntry.id)).where(JournalEntry.company_id == company_id)
    )).scalar()

    month_entries = (await db.execute(
        select(func.count(JournalEntry.id)).where(
            JournalEntry.company_id == company_id,
            JournalEntry.entry_date >= month_start,
        )
    )).scalar()

    totals = (await db.execute(
        select(
            func.coalesce(func.sum(JournalEntryLine.debit_amount), 0),
            func.coalesce(func.sum(JournalEntryLine.credit_amount), 0),
        )
        .select_from(JournalEntryLine)
        .join(JournalEntry, JournalEntry.id == JournalEntryLine.journal_entry_id)
        .where(JournalEntry.company_id == company_id)
    )).one()

    # Current period status
    current_period = (await db.execute(
        select(AccountingPeriod.status).where(
            AccountingPeriod.company_id == company_id,
            AccountingPeriod.start_date <= today,
            AccountingPeriod.end_date >= today,
        ).limit(1)
    )).scalar_one_or_none()

    return ResponseBase(data=GLAnalytics(
        total_accounts=total_accounts or 0, active_accounts=active_accounts or 0,
        total_journal_entries=total_entries or 0, entries_this_month=month_entries or 0,
        total_debit=Decimal(str(totals[0] or 0)),
        total_credit=Decimal(str(totals[1] or 0)),
        period_status=current_period,
    ))


# ── Exports ──────────────────────────────────────────────────────────────────

@router.get("/export/accounts")
async def export_chart_of_accounts(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_module("gl", "can_access")),
):
    rows = (await db.execute(
        select(GLAccount)
        .where(GLAccount.company_id == current_user.company_id)
        .order_by(GLAccount.code)
    )).scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ანგარიშთა გეგმა"
    ws.append(["კოდი", "სახელი", "ტიპი", "აქტიური", "აღწერა"])
    for a in rows:
        ws.append([a.code, a.name, a.account_type, "კი" if a.is_active else "არა", a.description or ""])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=chart_of_accounts.xlsx"})


@router.get("/export/journal")
async def export_journal_entries(
    date_from: str | None = None,
    date_to: str | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_module("gl", "can_access")),
):
    filters = [JournalEntry.company_id == current_user.company_id]
    if date_from:
        filters.append(JournalEntry.entry_date >= datetime.fromisoformat(date_from).date())
    if date_to:
        filters.append(JournalEntry.entry_date <= datetime.fromisoformat(date_to).date())

    entries = (await db.execute(
        select(JournalEntry).where(*filters)
        .options(selectinload(JournalEntry.lines))
        .order_by(JournalEntry.entry_date.desc())
    )).unique().scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ჟურნალი"
    ws.append(["ნომერი", "თარიღი", "აღწერა", "ტიპი", "ანგარიშის კოდი", "დებეტი", "კრედიტი"])
    for e in entries:
        for line in e.lines:
            ws.append([e.entry_number, str(e.entry_date), e.description, e.reference_type or "",
                       str(line.gl_account_id), float(line.debit_amount), float(line.credit_amount)])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=journal_entries.xlsx"})


@router.get("/export/trial-balance")
async def export_trial_balance(
    as_of_date: str | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_module("gl", "can_access")),
):
    target_date = date.fromisoformat(as_of_date) if as_of_date else date.today()
    rows = (await db.execute(
        select(
            GLAccount.code, GLAccount.name, GLAccount.account_type,
            func.coalesce(func.sum(JournalEntryLine.debit_amount), 0),
            func.coalesce(func.sum(JournalEntryLine.credit_amount), 0),
        )
        .select_from(GLAccount)
        .outerjoin(JournalEntryLine, JournalEntryLine.gl_account_id == GLAccount.id)
        .outerjoin(JournalEntry, JournalEntry.id == JournalEntryLine.journal_entry_id)
        .where(
            GLAccount.company_id == current_user.company_id,
            JournalEntry.entry_date <= target_date,
        )
        .group_by(GLAccount.id, GLAccount.code, GLAccount.name, GLAccount.account_type)
        .order_by(GLAccount.code)
    )).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "საცდელი ბალანსი"
    ws.append(["კოდი", "სახელი", "ტიპი", "დებეტი", "კრედიტი", "ბალანსი"])
    for code, name, atype, debit, credit in rows:
        balance = float(debit) - float(credit)
        ws.append([code, name, atype, float(debit), float(credit), balance])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=trial_balance.xlsx"})
