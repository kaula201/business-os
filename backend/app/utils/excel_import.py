"""Excel import utilities for Business OS."""
import io
from zipfile import BadZipFile
from datetime import datetime, date
from decimal import Decimal
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def parse_excel_upload(content: bytes, expected_columns: list[str]) -> list[dict]:
    """Parse an uploaded Excel file and return rows as dicts.

    Args:
        content: Raw bytes of the .xlsx file.
        expected_columns: Column names expected in the header row.

    Returns:
        List of dicts, one per data row (skipping the header).

    Raises:
        ValueError: If the file is empty, has no header, or required columns are missing.
    """
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("Excel ფაილი დაზიანებულია ან არასწორი ფორმატისაა") from exc
    ws = wb.active
    if ws is None or ws.max_row is None or ws.max_row < 2:
        raise ValueError("Excel ფაილი ცარიელია ან არ შეიცავს მონაცემებს")

    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c else "" for c in rows[0]]

    # Build column index map
    col_map = {}
    for i, h in enumerate(header):
        if h:
            col_map[h] = i

    # Validate required columns
    missing = [col for col in expected_columns if col.lower() not in col_map]
    if missing:
        raise ValueError(
            f"Excel ფაილში აკლია სვეტები: {', '.join(missing)}. "
            f"მოსალოდნელი სვეტები: {', '.join(expected_columns)}"
        )

    result = []
    for row in rows[1:]:
        if all(c is None for c in row):
            continue  # skip empty rows
        entry = {}
        for col_name in expected_columns:
            idx = col_map[col_name.lower()]
            entry[col_name] = row[idx] if idx < len(row) else None
        result.append(entry)

    return result


def safe_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def safe_float(value) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def safe_int(value) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def safe_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip()[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None
